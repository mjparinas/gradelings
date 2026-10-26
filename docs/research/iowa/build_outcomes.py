"""Builds outcomes.json from the Iowa Academic Standards (Iowa Department of Education; see fetch.sh and README.md).

Usage: python3 -I build_outcomes.py [source-dir]     (default: ./source)
Needs openpyxl. Reads the pdftotext copies (math.txt, ela.txt) and the Department's spreadsheets (science.xlsx,
social.xlsx) that fetch.sh saves.

Output: {"<grade>/<subject>": [{"idx", "kind", "content"}]} for Kindergarten ("k") to Grade 9 and the subjects math,
language, science and social. Grade 9 lists Algebra 1 (math), the Grades 9-10 band (English language arts), the high
school science standards and the Grades 9-12 social studies standards.

Codes, as the documents print them:
  math      grade, domain, cluster letter and number: "4.NBT.A.1", "K.CC.B.4"; Grade 9 cites the Algebra 1 standards
            with their course prefix: "A1.A-SSE.A.2", "A1.F-IF.B.4" (one Algebra 1 standard prints a part letter,
            "A1.A-SSE.A.1.a").
  language  strand, grade and number, printed after each standard: "RL.4.2", "RI.5.3", "RF.2.3", "W.4.1", "WF.4.1"
            (writing foundations), "SL.4.1", "L.4.4"; Grade 9 "RL.9-10.2".
  science   performance expectation codes from science.xlsx: "4-PS3-2"; middle school codes carry their grade
            ("6-LS1-1", "7-PS2-1", "8-LS4-1") and engineering is listed in every grade ("K-ETS1-1", "3-ETS1-2",
            "7-ETS1-4"); high school codes ("HS-ESS2-4") go under 9.
  social    "SS.K.6", "SS.4.12", "SS.8.24"; Grades 9-12 "SS.9-12.Gov.13", "SS.9-12.Econ.5", "SS.9-12.FL.4",
            "SS.9-12.Geo.3", "SS.9-12.US.4", "SS.9-12.WH.2" (inquiry "SS.9-12.3"). The spreadsheet ends most codes with a
            full stop ("SS.K.1."), which is dropped. It prints "SS.7.33" twice (the second is the Antarctic government
            standard) and has no "SS.7.37"; only the first SS.7.33 is listed.
"""

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import openpyxl

GRADES = ["k", "1", "2", "3", "4", "5", "6", "7", "8", "9"]
SUBJECTS = ["math", "language", "science", "social"]
OUT = OrderedDict((f"{g}/{s}", OrderedDict()) for g in GRADES for s in SUBJECTS)


def clean(text):
    text = text.replace("‐", "-").replace("–", "-").replace("’", "'").replace("“", '"').replace("”", '"')
    return re.sub(r"\s+", " ", text).strip()


def add(grade, subject, idx, content):
    entries = OUT[f"{grade}/{subject}"]
    if idx not in entries:
        entries[idx] = {"idx": idx, "kind": "standard", "content": clean(content)[:600]}


def grade_of(label):
    label = label.strip()
    return "k" if label.upper().startswith("K") else ("9" if label.startswith(("HS", "9")) else label[0])


# ---------------------------------------------------------------- math
RIGOR = re.compile(r"\s{2,}(Conceptual|Procedural|Application)\b.*$")
# The rigor and practice columns sometimes run into the standard's first line ("two- Conceptual Problem Solving digit").
COLUMNS = re.compile(r"\s(?:Conceptual|Procedural|Application)(?:\s+(?:Problem Solving|Communicating|Reasoning))*(?=\s)")
MATH_CODE = re.compile(r"^((?:K|[1-8])\.[A-Z]{1,3}\.[A-Z]\.\d+[a-z]?|A1\.[NAFGS]-[A-Z]{1,4}\.[A-Z]\.\d+(?:\.[a-z])?)\s+(\S.*)$")
DOMAIN = re.compile(r"^Domain:\s*((?:K|[1-8])\.[A-Z]{1,3}|A1\.[NAFGS]-[A-Z]{1,4}|[A-Z0-9.]*-?[A-Z]+)\b")


def math(src):
    """Grades K-8 and Algebra 1. Codes count only inside their own domain's section (the progressions tables repeat them)."""
    domain, current = None, None
    for line in (src / "math.txt").read_text(encoding="utf-8").split("\n"):
        d = DOMAIN.match(line.strip())
        if d:
            domain, current = d.group(1), None
            continue
        m = MATH_CODE.match(line)
        if m and domain and m.group(1).startswith(domain + "."):
            code = m.group(1)
            grade = "9" if code.startswith("A1.") else grade_of(code)
            current = (grade, code)
            add(grade, "math", code, COLUMNS.sub("", RIGOR.sub("", m.group(2))))
            continue
        if current and line.strip():
            entry = OUT[f"{current[0]}/math"][current[1]]
            if len(entry["content"]) < 400 and not line.startswith(" " * 40):
                entry["content"] = clean(entry["content"] + " " + re.split(r"\s{3,}", line.strip())[0])
        elif not line.strip():
            current = None


# ---------------------------------------------------------------- English language arts
# A code can break across a line ("(RL.9-" then "10.4)").
ELA_CODE = re.compile(r"\(\s*(RL|RI|RF|WF|W|SL|L)\.\s*(K|[1-8]|9-\s*10|11-\s*12)\s*\.\s*(\d{1,2})\s*\)")


def language(src):
    """The code follows each standard in parentheses; the standard's text is what comes before it since the last code."""
    text = (src / "ela.txt").read_text(encoding="utf-8")
    last = 0
    for m in ELA_CODE.finditer(text):
        chunk = text[last : m.start()]
        last = m.end()
        band = re.sub(r"\s+", "", m.group(2))
        if band == "11-12":
            continue
        grade = "9" if band == "9-10" else grade_of(band)
        # The standard starts at its number ("2. Determine a theme"); drop page furniture and the DOK note before it.
        chunk = re.sub(r"\(DOK[^)]*\)", " ", chunk)
        parts = re.split(r"\n\s*\d{1,2}\.\s+", chunk)
        if parts[-1].strip().startswith(("(Not applicable", "(Begins in grade")):
            continue
        add(grade, "language", f"{m.group(1)}.{band}.{m.group(3)}", parts[-1])


# ---------------------------------------------------------------- science and social studies (spreadsheets)
def science(src):
    wb = openpyxl.load_workbook(src / "science.xlsx", read_only=True)
    for ws in wb.worksheets[1:]:
        for row in ws.iter_rows(values_only=True):
            code = str(row[0] or "").strip()
            if re.match(r"^(K|[1-8]|HS)-[A-Z]{2,4}\d-\d+$", code):
                add(grade_of(code), "science", code, str(row[1] or ""))


def social(src):
    wb = openpyxl.load_workbook(src / "social.xlsx", read_only=True)
    for ws in wb.worksheets:
        for row in ws.iter_rows(values_only=True):
            cells = [str(c).strip() for c in row if c is not None]
            for i, c in enumerate(cells):
                m = re.match(r"^SS\.(K|[1-8]|9-12)\.((?:[A-Za-z]+\.)?\d+)\.?$", c)
                if m and i + 1 < len(cells):
                    add(grade_of(m.group(1)), "social", f"SS.{m.group(1)}.{m.group(2)}", cells[i + 1])
                    break


def main():
    src = Path(sys.argv[1] if len(sys.argv) > 1 else Path(__file__).parent / "source")
    math(src)
    language(src)
    science(src)
    social(src)
    data = {key: list(entries.values()) for key, entries in OUT.items()}
    for key, entries in data.items():
        if not entries:
            print("error: no standards for", key, file=sys.stderr)
            sys.exit(1)
    out = Path(__file__).parent / "outcomes.json"
    out.write_text(json.dumps(data, indent=1, ensure_ascii=False) + "\n", encoding="utf-8")
    print(f"{sum(len(e) for e in data.values())} standards written to {out.name}")


if __name__ == "__main__":
    main()
