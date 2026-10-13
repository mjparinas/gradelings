#!/usr/bin/env python3
"""Build outcomes.json from the OSPI files that fetch.sh downloads.

Usage: python3 build_outcomes.py <cache-dir>

Writes outcomes.json beside this script:
  {"<grade>/<subject>": [{"idx": code, "kind": "standard", "content": text}]}
Grades are "k" and "1".."9". Band standards are listed under every grade of
the band: middle school science (MS-) and social studies (6-8) under Grades 6
to 8; high school science (HS-), math (Algebra 1 and Integrated 1) and the
9-10 social studies and ELA standards under Grade 9.
Needs openpyxl and pypdf.
"""
import json
import re
import sys
from pathlib import Path

import openpyxl
from pypdf import PdfReader

GRADES = ["k", "1", "2", "3", "4", "5", "6", "7", "8", "9"]


def pdf_text(path: Path) -> str:
    return "\n".join(page.extract_text() or "" for page in PdfReader(path).pages)


def clean(text: str) -> str:
    return re.sub(r"\s+", " ", text).strip()


def grade_id(g: str) -> str:
    return "k" if g.upper() == "K" else g


def add(out: dict, grade: str, subject: str, idx: str, content: str) -> None:
    items = out.setdefault(f"{grade}/{subject}", [])
    if any(i["idx"] == idx for i in items):
        return
    items.append({"idx": idx, "kind": "standard", "content": content})


def math(cache: Path, out: dict) -> None:
    wb = openpyxl.load_workbook(cache / "math.xlsx", read_only=True)
    for row in list(wb["K-8 and Connections"].iter_rows(values_only=True))[3:]:
        code, grade, _domain, text = row[0], row[1], row[2], row[3]
        if code and text and str(code).startswith("M."):
            add(out, grade_id(str(grade)), "math", str(code).strip(), clean(str(text)))
    for row in list(wb["HS and Connections"].iter_rows(values_only=True))[3:]:
        code, course, text = row[1], row[2], row[4]
        if code and text and str(course) in ("ALG 1", "INT 1"):
            add(out, "9", "math", str(code).strip(), clean(str(text)))


ELA_CODE = re.compile(r"ELA\.(K|\d|9[-–]10)\.([A-Z]+)\.(\d+)")


def ela(cache: Path, out: dict) -> None:
    text = pdf_text(cache / "ela.pdf")
    matches = list(ELA_CODE.finditer(text))
    for n, m in enumerate(matches):
        # Only a code that starts a standard: at a line start, or after "PRIORITY:".
        line_start = text.rfind("\n", 0, m.start()) + 1
        before = text[line_start : m.start()].strip()
        if before not in ("", "PRIORITY:"):
            continue
        end = matches[n + 1].start() if n + 1 < len(matches) else len(text)
        body = text[m.end() : end]
        # The statement runs until its first lettered sub-item.
        body = re.split(r"\n\s*a\.\s", body, maxsplit=1)[0]
        body = clean(body.split("PRIORITY:")[0])
        if not body or body.startswith("Not in"):
            continue
        grade = m.group(1).replace("–", "-")
        grades = ["9"] if grade == "9-10" else [grade_id(grade)]
        # The PDF writes the band with an en dash; we cite it with a hyphen.
        for g in grades:
            add(out, g, "language", f"ELA.{grade}.{m.group(2)}.{m.group(3)}", body)


SCI_CODE = re.compile(r"^(K|[1-5]|MS|HS)-([A-Z]{2,4}\d)-(\d+)\.\s", re.M)


def science(cache: Path, out: dict) -> None:
    text = pdf_text(cache / "science-dci.pdf")
    matches = list(SCI_CODE.finditer(text))
    for n, m in enumerate(matches):
        end = matches[n + 1].start() if n + 1 < len(matches) else len(text)
        body = clean(re.split(r"\[Clarification|\[Assessment|\*", text[m.end() : end])[0])
        code = f"{m.group(1)}-{m.group(2)}-{m.group(3)}"
        band = m.group(1)
        grades = ["6", "7", "8"] if band == "MS" else ["9"] if band == "HS" else [grade_id(band)]
        for g in grades:
            add(out, g, "science", code, body)


SS_CODE = re.compile(r"\b((?:SSS|C|E|G|H)\d)\.(K|\d|6-8|9-10)\.(\d+)\b(?!-)")


def social(cache: Path, out: dict) -> None:
    for band in ("k-5", "6-8", "9-12"):
        for strand in ("ss-skills", "civics", "economics", "geography", "history"):
            text = pdf_text(cache / f"social-{band}-{strand}.pdf")
            matches = list(SS_CODE.finditer(text))
            for n, m in enumerate(matches):
                # A component starts its line; later mentions are cross-references.
                line_start = text.rfind("\n", 0, m.start()) + 1
                if text[line_start : m.start()].strip():
                    continue
                end = matches[n + 1].start() if n + 1 < len(matches) else len(text)
                body = text[m.end() : end]
                body = re.split(r"\n\s*(?:By the end of|\(Grade|\(Kinder|Since Time|S i n c e|Social Studies Learning Standards)", body)[0]
                # Join words the PDF hyphenated across lines ("associ- ated").
                body = re.sub(r"(\w)- (\w)", r"\1\2", clean(body.lstrip(":")))
                if not body:
                    continue
                grade = m.group(2)
                grades = ["6", "7", "8"] if grade == "6-8" else ["9"] if grade == "9-10" else [grade_id(grade)]
                for g in grades:
                    add(out, g, "social", m.group(0), body)


def main() -> None:
    cache = Path(sys.argv[1])
    out: dict = {}
    math(cache, out)
    ela(cache, out)
    science(cache, out)
    social(cache, out)
    ordered = {f"{g}/{s}": out.get(f"{g}/{s}", []) for g in GRADES for s in ("math", "language", "science", "social")}
    path = Path(__file__).with_name("outcomes.json")
    path.write_text(json.dumps(ordered, indent=1, ensure_ascii=False) + "\n")
    for key, items in ordered.items():
        print(key, len(items))


if __name__ == "__main__":
    main()
