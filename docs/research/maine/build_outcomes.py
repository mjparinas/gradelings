"""Builds outcomes.json from the Maine Department of Education's standards files (see fetch.sh and README.md).

Usage: python3 -I build_outcomes.py <cache-dir>
Needs pdfplumber and python-docx.

Output: {"<grade>/<subject>": [{"idx", "kind", "content", ...}]} for Kindergarten ("k") to Grade 9 and the subjects
math, language, science and social. Grade spans are listed under every grade of the span: Grades 6-8 under 6, 7 and 8,
and the 9-Diploma span under 9.
"""

import json
import re
import sys
from collections import OrderedDict
from pathlib import Path

import docx
import pdfplumber

GRADES = ["k", "1", "2", "3", "4", "5", "6", "7", "8", "9"]
MIDDLE = ["6", "7", "8"]


def fail(message):
    print("error:", message, file=sys.stderr)
    sys.exit(1)


def clean(text):
    text = re.sub(r"\s+", " ", text).strip()
    return text


def pdf_text(path):
    with pdfplumber.open(path) as pdf:
        return "\n".join((page.extract_text() or "") for page in pdf.pages)


# ---------------------------------------------------------------- shared wording
# Kindergarten to Grade 5 math and the science expectations are printed in side-by-side columns, so their text layer
# mixes neighbouring columns. Where the expectation is the same Common Core or NGSS one California records
# (docs/research/california/outcomes.json, read from CDE's single-column files), its wording is used instead. Codes
# always come from Maine's files.
CALIFORNIA = json.loads((Path(__file__).parent.parent / "california" / "outcomes.json").read_text())
CA_TEXT = {o["idx"]: o["content"] for entries in CALIFORNIA.values() for o in entries}


# Maine's one addition to the Common Core (Grade 1 money), read by hand from the childhood PDF's middle column.
MAINE_TEXT = {"1.MD.D.5": "Identify the coins and each corresponding value (e.g. penny, nickel, dime, and quarter)."}


def common_core_text(code):
    """California's wording of the same Common Core expectation ("4.NBT.A.1" is California's "4.NBT.1")."""
    if code in MAINE_TEXT:
        return MAINE_TEXT[code]
    m = re.fullmatch(r"([K1-8])\.([A-Z]+)\.[A-Z]\.(\d+)", code)
    if not m:
        return None
    return CA_TEXT.get(f"{m.group(1)}.{m.group(2)}.{m.group(3)}")


# ---------------------------------------------------------------- mathematics
# Maine's 2020 mathematics standards keep the Common Core performance expectations with their cluster letters
# ("4.NBT.A.1"). High school expectations are "HSN.RN.A.1" and belong to the 9-Diploma span (Grade 9 here). Lettered
# parts ("K.CC.B.4a", "HSN.VM.B4a") are part of their expectation and not recorded on their own.
MATH_CODE = re.compile(
    r"(?<![A-Za-z0-9.])(\(\+\)\s*)?((?:K|[1-8])\.(?:CC|OA|NBT|NF|MD|G|RP|NS|EE|F|SP)\.[A-Z]\.\d+|HS[NAFGS]\.[A-Z]{1,3}\.[A-Z]\.?\d+)([a-z]?)\s*:"
)


def math(cache):
    out = {g: OrderedDict() for g in GRADES}
    for name in ["math-childhood.pdf", "math-early-adolescence.pdf", "math-adolescence.pdf"]:
        # The PDFs' text layer sometimes splits a code ("8 .G.B.8").
        text = re.sub(r"(?<![0-9])([K1-8]) \.([A-Z])", r"\1.\2", pdf_text(cache / name))
        matches = list(MATH_CODE.finditer(text))
        for i, m in enumerate(matches):
            plus, code, part = m.group(1), m.group(2), m.group(3)
            if part:
                continue
            code = re.sub(r"^(HS[NAFGS]\.[A-Z]{1,3}\.[A-Z])(\d+)$", r"\1.\2", code)
            end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
            body = text[m.end() : end]
            body = re.split(r"\n(?:Strand|Standard|Performance|Expectations|\d+\n)", body)[0]
            body = clean(re.sub(r"Example:.*", "", body, flags=re.S))
            grade = "9" if code.startswith("HS") else code.split(".")[0].lower()
            if code in out[grade]:
                continue
            entry = {"idx": code, "kind": "expectation", "content": common_core_text(code) or body}
            if plus:
                entry["plus"] = True
            out[grade][code] = entry
    for g in GRADES:
        if len(out[g]) < 15:
            fail(f"only {len(out[g])} math expectations in grade {g}")
    return {g: list(v.values()) for g, v in out.items()}


# ---------------------------------------------------------------- docx tables (ELA and social studies)
def label_grades(label):
    """The grades a column heading covers, or None when it isn't a grade heading."""
    label = label.strip()
    if label == "Kindergarten":
        return ["k"]
    m = re.fullmatch(r"Grade (\d)", label)
    if m:
        return [m.group(1)]
    # Span headings read "Grades 6-8" or "Early Adolescence: Grades 6-8".
    if re.search(r"Grades? 6-8$", label):
        return MIDDLE
    if re.search(r"Grades? 9-Diploma$", label):
        return ["9"]
    return None


def unique_cells(row):
    """A row's cells, with horizontally merged cells counted once."""
    seen, cells = set(), []
    for cell in row.cells:
        if id(cell._tc) in seen:
            continue
        seen.add(id(cell._tc))
        cells.append(cell.text.strip())
    return cells


# ---------------------------------------------------------------- English language arts
# The 2020 ELA standards have no codes: four strands, each with numbered standards (Language 1-6, Speaking and
# Listening 1-4, Reading 1-12, Writing 1-3), and performance expectations per grade to Grade 5, then the spans 6-8 and
# 9-Diploma. A standard is cited as "Reading 4".
STRANDS = {"LANGUAGE": "Language", "READING": "Reading", "WRITING": "Writing", "SPEAKING AND LISTENING": "Speaking and Listening"}
# Foundational reading (Reading 1-3) ends at Grade 5; the spans after it only point back to earlier grades.
SKIP_TEXT = re.compile(r"^(n/?a|not applicable|\(?begins in|\(?see grade|some learners may enter)", re.I)


def ela(cache):
    document = docx.Document(cache / "ela.docx")
    out = {g: OrderedDict() for g in GRADES}
    strand = number = statement = None
    for table in document.tables:
        grades = None
        for row in table.rows:
            cells = unique_cells(row)
            head, rest = cells[0], cells[1:]
            if head == "Strand":
                key = rest[0].split(":")[0].split("/")[0].strip().upper()
                if key not in STRANDS:
                    fail(f"unknown ELA strand {rest[0]!r}")
                strand = STRANDS[key]
                continue
            m = re.fullmatch(r"Standard (\d+)", head)
            if m:
                number, statement = int(m.group(1)), clean(rest[0])
                continue
            labels = [label_grades(c) for c in rest]
            if rest and all(labels):
                grades = labels
                continue
            if head.startswith("Performance") and grades and len(rest) == len(grades):
                for gs, text in zip(grades, rest):
                    text = clean(text)
                    if not text or SKIP_TEXT.match(text):
                        continue
                    for g in gs:
                        idx = f"{strand} {number}"
                        entry = out[g].setdefault(idx, {"idx": idx, "kind": "standard", "content": statement, "expectations": []})
                        if text not in entry["expectations"]:
                            entry["expectations"].append(text)
                grades = None
    expected = {"Language": 6, "Speaking and Listening": 4, "Reading": 12, "Writing": 3}
    found = {s: {int(i.split()[-1]) for g in GRADES for i in out[g] if i.rsplit(" ", 1)[0] == s} for s in expected}
    for s, n in expected.items():
        if found[s] != set(range(1, n + 1)):
            fail(f"ELA {s} standards found: {sorted(found[s])}")
    return {g: list(v.values()) for g, v in out.items()}


# ---------------------------------------------------------------- science
# The 2019 Science and Engineering standards are the Next Generation Science Standards' performance expectations with
# their codes ("4-PS3-2", "MS-LS1-1", "HS-ESS2-1"). Middle school is one 6-8 span (listed under 6, 7 and 8); high
# school is the 9-Diploma span (Grade 9 here). K-2-ETS1 is under Kindergarten to Grade 2 and 3-5-ETS1 under 3 to 5.
SCI_CODE = re.compile(r"(?<![A-Za-z0-9-])((?:K-2|3-5|K|[1-5]|MS|HS)-(?:PS|LS|ESS|ETS)\d-\d+)(?![0-9])")


def science(cache):
    text = pdf_text(cache / "science.pdf")
    out = {g: OrderedDict() for g in GRADES}
    matches = list(SCI_CODE.finditer(text))
    for i, m in enumerate(matches):
        code = m.group(1)
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        body = re.split(r"Further explanation|Expectations\s*\n|\n(?:Strand|Standard|Performance)", text[m.end() : end])[0]
        body = clean(body.replace("Expectations", ""))
        band = code.rsplit("-", 2)[0]
        grades = {"K-2": ["k", "1", "2"], "3-5": ["3", "4", "5"], "MS": MIDDLE, "HS": ["9"], "K": ["k"]}.get(band, [band])
        for g in grades:
            out[g].setdefault(code, {"idx": code, "kind": "expectation", "ngss": code, "content": CA_TEXT.get(code) or body})
    for g in GRADES:
        if len(out[g]) < 8:
            fail(f"only {len(out[g])} science expectations in grade {g}")
    return {g: list(v.values()) for g, v in out.items()}


# ---------------------------------------------------------------- social studies
# The 2019 Social Studies standards have four strands (Civics & Government, Personal Finance & Economics, Geography,
# History). Each grade from Kindergarten to Grade 5, and the spans 6-8 and 9-Diploma, has named performance
# expectations: "Civics & Government 1", "Personal Finance", "Economics", "Global Connections", "Geography 2",
# "History 3". In the spans each has foundational (F1, F2) and developmental (D1, D2) parts, kept as `parts`.
SOCIAL_NAME = re.compile(r"^(Civics & Government \d|Personal Finance|Economics|Global Connections|Geography \d|History \d)\s*:\s*(.*)$")


def social(cache):
    document = docx.Document(cache / "social.docx")
    out = {g: OrderedDict() for g in GRADES}
    for table in document.tables:
        grades = None
        current = None
        for row in table.rows:
            cells = unique_cells(row)
            head, rest = cells[0], cells[1:]
            labels = [label_grades(c) for c in rest]
            if rest and all(labels):
                grades = labels
                continue
            # Later rows of a span table leave the first column empty.
            if not (head.startswith("Performance") or head == "") or not grades:
                continue
            # A column per grade (or several columns of one span: foundational, then developmental).
            columns = [(gs, text) for gs, text in zip(grades, rest)]
            if len({tuple(gs) for gs, _ in columns}) == 1 and columns[0][0] in (MIDDLE, ["9"]):
                gs = columns[0][0]
                first = clean(columns[0][1])
                m = SOCIAL_NAME.match(first)
                if m:
                    current = m.group(1)
                    for g in gs:
                        out[g].setdefault(current, {"idx": current, "kind": "expectation", "content": clean(m.group(2)), "parts": []})
                    continue
                # The span heading is one merged cell over a foundational (F) and a developmental (D) column.
                for text in rest:
                    for line in text.split("\n"):
                        pm = re.match(r"\((F|D)\s*(\d)\)\s*(.*)", line.strip())
                        if pm and current:
                            for g in gs:
                                part = f"{pm.group(1)}{pm.group(2)} {clean(pm.group(3))}"
                                if part not in out[g][current]["parts"]:
                                    out[g][current]["parts"].append(part)
                continue
            for gs, text in columns:
                for line in text.split("\n"):
                    m = SOCIAL_NAME.match(line.strip())
                    if not m:
                        continue
                    for g in gs:
                        out[g].setdefault(m.group(1), {"idx": m.group(1), "kind": "expectation", "content": clean(m.group(2))})
    for g in GRADES:
        names = list(out[g])
        if not any(n.startswith("History") for n in names) or not any(n.startswith("Geography") for n in names):
            fail(f"grade {g} social studies is missing a strand: {names}")
    return {g: list(v.values()) for g, v in out.items()}


def main():
    if len(sys.argv) != 2:
        fail("usage: build_outcomes.py <cache-dir>")
    cache = Path(sys.argv[1])
    subjects = {"math": math(cache), "language": ela(cache), "science": science(cache), "social": social(cache)}
    result = {f"{g}/{s}": subjects[s][g] for g in GRADES for s in subjects}
    path = Path(__file__).with_name("outcomes.json")
    path.write_text(json.dumps(result, indent=1, ensure_ascii=False) + "\n")
    for s in subjects:
        print(s, " ".join(f"{g}:{len(subjects[s][g])}" for g in GRADES))


if __name__ == "__main__":
    main()
